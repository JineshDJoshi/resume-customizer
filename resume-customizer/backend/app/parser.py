"""Extract plain text from uploaded PDF / DOCX resumes."""
import io
import zipfile

import pdfplumber
from docx import Document

MAX_PDF_PAGES = 10
MIN_TEXT_CHARS = 80


class ResumeParseError(Exception):
    def __init__(self, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.status_code = status_code


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        if not data.startswith(b"%PDF"):
            raise ResumeParseError("This file has a .pdf extension but is not a valid PDF.", 415)
        text = _pdf_text(data)
    elif ext == "docx":
        if not data.startswith(b"PK"):
            raise ResumeParseError("This file has a .docx extension but is not a valid DOCX.", 415)
        text = _docx_text(data)
    else:
        raise ResumeParseError("Unsupported file type. Upload a PDF or DOCX resume.", 415)

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeParseError(
            "We couldn't read any text from this resume. It may be a scanned image — "
            "export a text-based PDF or upload a DOCX instead."
        )
    return text


def _pdf_text(data: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(page.extract_text() or "") for page in pdf.pages[:MAX_PDF_PAGES]]
    except Exception as exc:  # pdfplumber/pdfminer raise many different types
        raise ResumeParseError(
            "This PDF could not be opened. It may be corrupted or password-protected."
        ) from exc
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ResumeParseError("This DOCX could not be opened. It may be corrupted.") from exc
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                if cell.text and cell.text not in cells:
                    cells.append(cell.text)
            lines.append(" | ".join(cells))
    return "\n".join(lines)
